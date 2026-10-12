"""DOCX 渲染器 —— python-docx 直接构建。

ADR-P8-001：**不用 docxtpl 模板填充**。

## 不嵌字体

DOCX **不嵌入字体文件**，中文由打开文档的环境决定 —— python-docx 也做不到。
不要试图塞字体：中文字库 10MB 起步，文件体积爆炸而毫无作用。
文档里说明「打开需要中文字体环境」即可。

## 页码

`{{page}}` 用 Word 的 `PAGE` **域**实现，而不是留字面量或留空。
留字面量会在对外报告上印出一串 `{{page}}`；留空则是「第 页」。
域是 Word/其他 OOXML 阅读器的标准做法，python-docx 没有高层 API，
故直接写 `w:fldChar` XML。
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.services.report.blocks import (
    Block,
    HeadingBlock,
    KeyValueBlock,
    PageBreakBlock,
    TableBlock,
    _cell_text,
    assumed_header,
    fill_placeholders,
)
from app.services.report.style import ReportStyle

_PAGESIZES_MM = {"A4": (210.0, 297.0), "A3": (297.0, 420.0), "Letter": (215.9, 279.4)}


def _add_page_field(paragraph) -> None:
    """往段落里插一个 Word `PAGE` 域（真正的页码，会随页自动更新）。"""
    run = paragraph.add_run()
    begin = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = run._r.makeelement(qn("w:instrText"), {})
    instr.text = " PAGE "
    end = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    for el in (begin, instr, end):
        run._r.append(el)


def _write_hf_parts(hf, parts, context: dict, content_width) -> None:
    """页眉/页脚**单段 + 制表位**：左 \| 中 \| 右 同一行。

    多段写法每段独占一行（页眉 3 行、页脚 2 行），白占页眉高度。Word 惯例是
    一个段落 + 中/右制表位。tab 位置按**当前版面正文宽**算，不能吃模板默认值
    —— 默认 Header 样式的制表位是按 Letter 纸定的，A4/Landscape 下右对齐会
    错位。section 自带的空段（paragraphs[0]）直接复用，避免开头多一空行。
    """
    from docx.enum.text import WD_TAB_ALIGNMENT

    p = hf.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(
        int(content_width / 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(
        int(content_width), WD_TAB_ALIGNMENT.RIGHT)

    first = True
    for _align, part in parts:
        if not part:
            continue
        if not first:
            p.add_run("\t")
        first = False
        before, _, after = fill_placeholders(part, context).partition("{{page}}")
        if before:
            p.add_run(before)
        if _:
            _add_page_field(p)
        if after:
            p.add_run(after)


def render_docx(*, title: str, blocks: list[Block], style: ReportStyle,
                context: dict) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    w, h = _PAGESIZES_MM[style.page_size]
    if style.orientation == "landscape":
        w, h = h, w
        sec.orientation = WD_ORIENT.LANDSCAPE
    # ⚠️ 只设 orientation 不会让 Word 重排页面，必须**同时**交换宽高
    sec.page_width, sec.page_height = Mm(w), Mm(h)
    m = style.margins
    sec.top_margin, sec.bottom_margin = Mm(m.top), Mm(m.bottom)
    sec.left_margin, sec.right_margin = Mm(m.left), Mm(m.right)

    # 左 | 中 | 右 同一行（tab 分隔），tab 位按正文宽算。footer.left 是
    # doc_no、footer.right 是「第 N 页」—— 缺一段就是悄悄丢内容。
    content_w = sec.page_width - sec.left_margin - sec.right_margin
    _write_hf_parts(sec.header, (
        (None, style.header.left), (None, style.header.center),
        (None, style.header.right)), context, content_w)
    _write_hf_parts(sec.footer, (
        (None, style.footer.left), (None, style.footer.center),
        (None, style.footer.right)), context, content_w)

    doc.add_heading(title, level=1)
    mc = style.table.max_cell_chars
    for b in blocks:
        if isinstance(b, HeadingBlock):
            doc.add_heading(b.text, level=min(b.level, 4))
        elif isinstance(b, PageBreakBlock):
            doc.add_page_break()
        elif isinstance(b, KeyValueBlock):
            t = doc.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for k, v in b.rows:
                cells = t.add_row().cells
                cells[0].text = str(k)
                cells[1].text = _cell_text(v, mc)
        elif isinstance(b, TableBlock):
            t = doc.add_table(rows=1, cols=len(b.headers))
            t.style = "Table Grid"
            for i, head in enumerate(b.headers):
                t.rows[0].cells[i].text = assumed_header(
                    head, i in b.assumed_cols, style.assumption_mark)
            for row in b.rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = _cell_text(v, mc)
                    if i in b.assumed_cols:
                        _color_run(cells[i], style.assumption_color)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _color_run(cell, color: str) -> None:
    """把单元格里所有 run 染成指定色（SPEC §3.2.1(4) 假设字段标记）。

    `cell.text = ...` 会重置段落，所以染色必须在赋值**之后**做。
    """
    rgb = RGBColor.from_string(color.lstrip("#"))
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = rgb
            run.font.size = Pt(10)
