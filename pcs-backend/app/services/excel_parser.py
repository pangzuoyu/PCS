"""P3.2 SIM-5：Excel 批量导入解析器（spec 第二部分）。

按 spec §2.2 落地：
- 双 Sheet 结构：物流列表 + 组分组成（按 Stream Name 关联）
- 单位：temp °C → K（+273.15）；press kPa → Pa（×1000）
- Mole Fraction 与 Mass Flow 至少填一列；两者都填时以 Mole Fraction 为准
- 组分名别名表（H₂O→WATER 等，spec 附录 A）
- 错误兜底：缺列、组分行两列空、损坏文件
- 组分键为名称（非 CAS），CAS 转换在 SIM-10 commit 阶段完成

输出 ExcelParseResult：streams（list[ParsedStream]）+ warnings + errors。
下游：SIM-10 导入预览 commit 复用本 parser，冲突检测走 SIM-7。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from app.services.property_completion import ParsedStream

_C_TO_K_OFFSET = 273.15
_KPA_TO_PA = 1000.0

# Sheet 1（物流列表）必填列（spec §2.2）
REQUIRED_STREAM_COLS = (
    "Stream Name",
    "Temp (°C)",
    "Pressure (kPa)",
    "Phase",
    "Total Mass Flow (kg/h)",
    "Total Molar Flow (kmol/h)",
    "Description",
)

# Sheet 2（组分组成）列
COMPONENT_COL_STREAM = "Stream Name"
COMPONENT_COL_NAME = "Component Name"
COMPONENT_COL_MOLE = "Mole Fraction"
COMPONENT_COL_MASS = "Mass Flow (kg/h)"

# Sheet 名（中文/英文 fallback）
SHEET1_NAMES = ("物流列表", "Streams", "Stream List")
SHEET2_NAMES = ("组分组成", "Components", "Compositions")

# spec 附录 A 名称别名（最小集：H₂O→WATER）
# 后续 P3.2 SIM-10 阶段可扩展（CH4/METHANE、NH3/AMMONIA 等）
COMPONENT_ALIASES: dict[str, str] = {
    "H₂O": "WATER",
    "H2O": "WATER",
    "水": "WATER",
}


@dataclass
class ExcelParseResult:
    """Excel 解析结果。"""

    streams: list[ParsedStream] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _norm_header(s: object) -> str:
    """列名规整：去 BOM + 前后空白。"""
    return str(s or "").replace("﻿", "").strip()


def _find_sheet(wb, names: tuple[str, ...]) -> str | None:
    for n in names:
        if n in wb.sheetnames:
            return n
    return None


def _row_to_dict(header_row: tuple, data_row: tuple) -> dict[str, object]:
    """行 → {列名: 值}（header 为表头）。"""
    out: dict[str, object] = {}
    for i, h in enumerate(header_row):
        h_norm = _norm_header(h)
        if not h_norm:
            continue
        out[h_norm] = data_row[i] if i < len(data_row) else None
    return out


def _parse_stream_meta(row: dict[str, object]) -> ParsedStream:
    """Sheet 1 行 → ParsedStream（不填 composition，由 Sheet 2 后置合并）。"""
    tag = str(row.get("Stream Name") or "").strip()
    if not tag:
        # 跳过空行（保留 caller 错误记录能力）
        raise ValueError("empty stream name")
    # temp °C → K
    t_raw = row.get("Temp (°C)")
    temperature_k = float(t_raw) + _C_TO_K_OFFSET if t_raw not in (None, "") else None
    # press kPa → Pa
    p_raw = row.get("Pressure (kPa)")
    pressure_pa = float(p_raw) * _KPA_TO_PA if p_raw not in (None, "") else None
    # phase 大写
    phase = str(row.get("Phase") or "").strip().upper() or None
    # flows
    m_raw = row.get("Total Mass Flow (kg/h)")
    mass_flow = float(m_raw) if m_raw not in (None, "") else None
    n_raw = row.get("Total Molar Flow (kmol/h)")
    molar_flow = float(n_raw) if n_raw not in (None, "") else None
    return ParsedStream(
        tag=tag,
        temperature_k=temperature_k,
        pressure_pa=pressure_pa,
        phase=phase,
        mass_flow_kg_h=mass_flow,
        molar_flow_kmol_h=molar_flow,
    )


def _alias(name: str) -> str:
    """组分名别名规整（spec 附录 A）。"""
    return COMPONENT_ALIASES.get(name, name)


def parse_excel_file(path: Path | str) -> ExcelParseResult:
    """Excel 双 Sheet 解析入口（spec §2.2 + §2.3 解析段）。

    Args:
        path: .xlsx 文件路径

    Returns:
        ExcelParseResult：streams + warnings + errors

    Raises:
        ValueError: 缺 Sheet / 缺必填列 / 损坏文件
        FileNotFoundError: 路径不存在
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"无法读取 Excel 文件 {path}: {e}") from e

    try:
        sheet1 = _find_sheet(wb, SHEET1_NAMES)
        sheet2 = _find_sheet(wb, SHEET2_NAMES)
        if sheet1 is None:
            raise ValueError(
                f"Excel 缺 Sheet 1（{SHEET1_NAMES} 任一）: {path}"
            )
        # 严格 spec：必须双 Sheet（无 Sheet 2 视为不完整）
        # 但若用户只想录元数据，可在 errors 收集行级错
        # 此处按 spec 强校验：缺 Sheet 2 抛错
        if sheet2 is None:
            raise ValueError(
                f"Excel 缺 Sheet 2（{SHEET2_NAMES} 任一）: {path}"
            )

        ws1 = wb[sheet1]
        ws2 = wb[sheet2]

        # Sheet 1: 解析流列表
        rows1 = list(ws1.iter_rows(values_only=True))
        if len(rows1) < 2:
            return ExcelParseResult(
                warnings=[f"Sheet 1 无数据行: {path}"],
            )
        header1 = rows1[0]
        # 列名校验
        headers_norm = {_norm_header(h) for h in header1}
        missing = [c for c in REQUIRED_STREAM_COLS if c not in headers_norm]
        if missing:
            raise ValueError(
                f"Sheet 1 missing required column(s): {missing}"
            )

        # 1st pass：构造流（无 composition）
        result = ExcelParseResult()
        tag_to_stream: dict[str, ParsedStream] = {}
        for i, row in enumerate(rows1[1:], start=2):  # row 2 起
            if not any(c not in (None, "") for c in row):
                continue  # 跳过空行
            try:
                row_dict = _row_to_dict(header1, row)
                stream = _parse_stream_meta(row_dict)
                tag_to_stream[stream.tag] = stream
            except ValueError as e:
                result.errors.append(f"Sheet1 row {i}: {e}")

        # Sheet 2: 解析组分 → 合并到 composition
        rows2 = list(ws2.iter_rows(values_only=True))
        comp_header = rows2[0] if rows2 else ()
        comp_header_norm = [_norm_header(h) for h in comp_header]
        try:
            idx_stream = comp_header_norm.index(COMPONENT_COL_STREAM)
            idx_name = comp_header_norm.index(COMPONENT_COL_NAME)
            idx_mole = comp_header_norm.index(COMPONENT_COL_MOLE)
            idx_mass = comp_header_norm.index(COMPONENT_COL_MASS)
        except ValueError as e:
            raise ValueError(f"Sheet 2 列缺失: {e}") from e

        for i, row in enumerate(rows2[1:], start=2):
            if not any(c not in (None, "") for c in row):
                continue
            tag = str(row[idx_stream] or "").strip()
            name_raw = str(row[idx_name] or "").strip()
            mole = row[idx_mole]
            mass = row[idx_mass]
            if not tag or not name_raw:
                result.errors.append(
                    f"Sheet2 row {i}: 缺 stream_name 或 component_name"
                )
                continue
            if mole in (None, "") and mass in (None, ""):
                result.errors.append(
                    f"Sheet2 row {i}: Mole Fraction 与 Mass Flow 至少填一列"
                    f"（stream='{tag}' component='{name_raw}'）"
                )
                continue
            # spec 规则：两者都填时以 Mole Fraction 为准
            if mole not in (None, ""):
                value = float(mole)
            else:
                value = float(mass)  # 后续由 SIM-10 转换；此处仅占位
            name = _alias(name_raw)
            if tag not in tag_to_stream:
                result.errors.append(
                    f"Sheet2 row {i}: stream_name='{tag}' 不在 Sheet 1"
                )
                continue
            stream = tag_to_stream[tag]
            if stream.composition is None:
                stream = ParsedStream(
                    tag=stream.tag,
                    cas=stream.cas,
                    temperature_k=stream.temperature_k,
                    pressure_pa=stream.pressure_pa,
                    phase=stream.phase,
                    mass_flow_kg_h=stream.mass_flow_kg_h,
                    molar_flow_kmol_h=stream.molar_flow_kmol_h,
                    molecular_weight=stream.molecular_weight,
                    composition={name: value},
                    vapor_composition=stream.vapor_composition,
                    liquid_composition=stream.liquid_composition,
                )
                tag_to_stream[tag] = stream
            else:
                # dataclass(frozen=True) → 重建
                new_comp = dict(stream.composition)
                if name in new_comp:
                    result.warnings.append(
                        f"Sheet2 row {i}: 组分 '{name}' 在 stream='{tag}' 重复，"
                        "后者覆盖前者"
                    )
                new_comp[name] = value
                tag_to_stream[tag] = ParsedStream(
                    tag=stream.tag,
                    cas=stream.cas,
                    temperature_k=stream.temperature_k,
                    pressure_pa=stream.pressure_pa,
                    phase=stream.phase,
                    mass_flow_kg_h=stream.mass_flow_kg_h,
                    molar_flow_kmol_h=stream.molar_flow_kmol_h,
                    molecular_weight=stream.molecular_weight,
                    composition=new_comp,
                    vapor_composition=stream.vapor_composition,
                    liquid_composition=stream.liquid_composition,
                )

        result.streams = list(tag_to_stream.values())
        # 缺组分的 stream WARN
        for s in result.streams:
            if s.composition is None:
                result.warnings.append(
                    f"stream '{s.tag}' 无组分（Sheet 2 缺该 stream_name），"
                    "需手工补"
                )

        return result
    finally:
        wb.close()


__all__ = [
    "ExcelParseResult",
    "parse_excel_file",
    "COMPONENT_ALIASES",
    "REQUIRED_STREAM_COLS",
]
