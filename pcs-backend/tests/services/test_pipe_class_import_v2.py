"""SUP-002 PC-5：管道等级 Excel 双 Sheet + import_id 暂存（V1.4 §0.6/EXCEL-01）。

覆盖：
  1. test_build_import_template_12_columns            — 模板仍走 1.9.6 12 列
  2. test_preview_valid_excel_returns_import_id        — 合法 Excel → preview + UUID
  3. test_preview_invalid_dn_returns_pc_v02_errors     — dn_series 缺 → PC-V02 ERROR
  4. test_preview_pressure_out_of_range_returns_errors — dp=43 → PC-V05 ERROR
  5. test_preview_persists_row_to_db                   — preview 后 DB 行存在
  6. test_get_preview_roundtrip_matches                — preview → get_preview 内容一致
  7. test_get_preview_expired_raises_410               — expires_at 改过去 → 410
  8. test_get_preview_unknown_raises_404              — 未知 UUID → 404
  9. test_commit_import_writes_all_rows                — preview → commit → PipeClass 行存在
 10. test_commit_import_rejects_when_errors            — preview 含 ERROR → 422, DB 无新行
 11. test_commit_import_rejects_double_consume         — 第二次同 import_id → 409 IMPORT_CONSUMED
 12. test_import_validation_engine_used               — 实际调用 validator 22 条规则

契约要点（V1.4 §0.6/§三、#6）：
- Sheet1 12 列（或可选 13 列 BaseMaterial）；Sheet2 可选 Sch系列 长表覆盖
- preview 写 pipe_class_import_previews 返 import_id，TTL 24h
- commit_import 按 import_id 重放；二次同 import_id 拒绝（consumed_at）
"""
from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta

import openpyxl
import pytest
from openpyxl import Workbook
from sqlalchemy import update

from app.models.config_domain import PipeClass, PipeClassImportPreview
from app.services.exceptions import PcsError
from app.services.pipe_class_import_service import (
    ImportPreview,
    PipeClassImportService,
)
from app.services.pipe_class_validator import (
    PipeClassValidator,
    Severity,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

SHEET1_12 = [
    "ClassID", "ClassName", "MaterialStandard", "DesignPressure", "DesignTemp",
    "CorrosionAllowance", "FlangeClass", "FittingType", "AllowableStressTable",
    "BranchTable", "Source", "Version",
]


def _make_sheet1_row(
    *,
    class_id: str = "A1",
    class_name: str = "等级A1",
    material_standard: str = "GB/T 8163",
    base_material: str | None = None,
    design_pressure: float = 2.5,
    design_temperature: float = 200.0,
    corrosion_allowance: float = 1.5,
    flange_class: str = "PN25",
    fitting_type: str = "对焊",
    allowable_stress_table: str = "ASME_B31_3_TABLE_A1",
    branch_table: str = "BRANCH_TABLE_01",
    source: str = "COMPANY_STD",
    version: str = "PC-V1.0",
) -> list:
    """生成 12 或 13 列 Sheet1 数据行（base_material 决定列数）。"""
    if base_material is None:
        return [
            class_id, class_name, material_standard,
            design_pressure, design_temperature, corrosion_allowance,
            flange_class, fitting_type, allowable_stress_table,
            branch_table, source, version,
        ]
    # 13 列：BaseMaterial 插在 MaterialStandard 之后
    return [
        class_id, class_name, material_standard, base_material,
        design_pressure, design_temperature, corrosion_allowance,
        flange_class, fitting_type, allowable_stress_table,
        branch_table, source, version,
    ]


def _xlsx(
    *,
    rows: list[list] | None = None,
    headers: list[str] | None = None,
    sheet2: list[list] | None = None,
    sheet1_name: str = "等级列表",
    sheet2_name: str = "Sch系列",
    include_sheet2: bool = True,
) -> bytes:
    """构造测试 Excel。rows=None 时空表；sheet2 给定时附带 Sch系列 sheet。"""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = sheet1_name
    ws1.append(headers or SHEET1_12)
    for r in rows or []:
        ws1.append(r)
    if include_sheet2 and sheet2 is not None:
        ws2 = wb.create_sheet(sheet2_name)
        ws2.append(["ClassID", "DN", "Sch列表"])
        for r in sheet2:
            ws2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _valid_sheet2(class_id: str = "A1") -> list[list]:
    """最小合法 Sheet2：DN15 / DN50 → STD/40。"""
    return [
        [class_id, 15, "STD"],
        [class_id, 50, "40,STD"],
    ]


# ---------------------------------------------------------------------------
# 1. 模板下载
# ---------------------------------------------------------------------------


def test_build_import_template_12_columns():
    """build_import_template 返 12 列（V1.4 EXCEL-01 锁 1.9.6 模板）。"""
    from app.services.pipe_class_import_service import (
        SHEET1_HEADERS_12,
        PipeClassImportService,
    )

    assert len(SHEET1_HEADERS_12) == 12
    assert SHEET1_HEADERS_12[0] == "ClassID"

    buf = PipeClassImportService.build_import_template()
    wb = openpyxl.load_workbook(io.BytesIO(buf), read_only=True)
    ws = wb.active
    header_row = next(ws.iter_rows(values_only=True))
    assert list(header_row) == SHEET1_HEADERS_12


# ---------------------------------------------------------------------------
# 2. preview 主路径
# ---------------------------------------------------------------------------


async def test_preview_valid_excel_returns_import_id(db):
    """合法 Excel → preview + 非空 UUID。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row(class_id="A1")],
        sheet2=_valid_sheet2("A1"),
    )
    svc = PipeClassImportService(db)
    preview, import_id = await svc.preview(xlsx_bytes)
    assert isinstance(import_id, uuid.UUID) and import_id.version == 4
    assert isinstance(preview, ImportPreview)
    assert len(preview.valid) == 1
    assert preview.errors == []
    assert preview.valid[0].class_id == "A1"


# ---------------------------------------------------------------------------
# 3/4. 校验错误
# ---------------------------------------------------------------------------


async def test_preview_invalid_dn_returns_pc_v02_errors(db):
    """无 Sheet2 → dn_series 缺 → PC-V02 ERROR。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row(class_id="A1")],
        sheet2=[],
        include_sheet2=False,  # 无 Sch系列 sheet
    )
    svc = PipeClassImportService(db)
    preview, _ = await svc.preview(xlsx_bytes)
    assert preview.errors, "无 Sheet2 应触发 PC-V02"
    assert any(
        r.rule_id == "PC-V02" and r.severity == Severity.ERROR
        for r in preview.errors
    )
    assert preview.valid == []


async def test_preview_pressure_out_of_range_returns_errors(db):
    """dp=43 → PC-V05 ERROR。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row(class_id="A1", design_pressure=43.0)],
        sheet2=_valid_sheet2("A1"),
    )
    svc = PipeClassImportService(db)
    preview, _ = await svc.preview(xlsx_bytes)
    assert any(
        r.rule_id == "PC-V05" and r.severity == Severity.ERROR
        for r in preview.errors
    )


# ---------------------------------------------------------------------------
# 5. preview DB 持久化
# ---------------------------------------------------------------------------


async def test_preview_persists_row_to_db(db):
    """preview 后 pipe_class_import_previews 表应有行。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row()],
        sheet2=_valid_sheet2(),
    )
    svc = PipeClassImportService(db)
    _, import_id = await svc.preview(xlsx_bytes)
    row = await db.get(PipeClassImportPreview, import_id)
    assert row is not None
    assert row.import_id == import_id
    assert row.file_bytes == xlsx_bytes
    assert row.error_count == 0
    # SQLite 存 TIMESTAMP 丢 tz，比较时统一为 UTC
    expires = row.expires_at
    if expires.tzinfo is None:
        expires = expires.replace(tzinfo=UTC)
    assert expires > datetime.now(UTC)


# ---------------------------------------------------------------------------
# 6/7/8. get_preview
# ---------------------------------------------------------------------------


async def test_get_preview_roundtrip_matches(db):
    """preview → get_preview → 内容一致。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row(class_id="A1"), _make_sheet1_row(class_id="A2")],
        sheet2=_valid_sheet2("A1") + _valid_sheet2("A2"),
    )
    svc = PipeClassImportService(db)
    preview, import_id = await svc.preview(xlsx_bytes)
    fetched = await svc.get_preview(import_id)
    assert fetched.valid and not fetched.errors
    class_ids = {r.class_id for r in fetched.valid}
    assert class_ids == {"A1", "A2"}


async def test_get_preview_expired_raises_410(db):
    """expires_at 改过去 → 410 IMPORT_EXPIRED。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row()],
        sheet2=_valid_sheet2(),
    )
    svc = PipeClassImportService(db)
    _, import_id = await svc.preview(xlsx_bytes)
    # 手工回拨 expires_at
    await db.execute(
        update(PipeClassImportPreview)
        .where(PipeClassImportPreview.import_id == import_id)
        .values(expires_at=datetime.now(UTC) - timedelta(hours=1))
    )
    await db.commit()
    with pytest.raises(PcsError) as ei:
        await svc.get_preview(import_id)
    assert ei.value.code == "IMPORT_EXPIRED"
    assert ei.value.status == 410


async def test_get_preview_unknown_raises_404(db):
    """未知 UUID → 404 IMPORT_NOT_FOUND。"""
    svc = PipeClassImportService(db)
    with pytest.raises(PcsError) as ei:
        await svc.get_preview(uuid.uuid4())
    assert ei.value.code == "IMPORT_NOT_FOUND"
    assert ei.value.status == 404


# ---------------------------------------------------------------------------
# 9/10/11. commit_import
# ---------------------------------------------------------------------------


async def test_commit_import_writes_all_rows(db):
    """preview → commit → PipeClass 行写入。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row(class_id="A1"), _make_sheet1_row(class_id="A2")],
        sheet2=_valid_sheet2("A1") + _valid_sheet2("A2"),
    )
    svc = PipeClassImportService(db)
    _, import_id = await svc.preview(xlsx_bytes)
    inserted = await svc.commit_import(import_id)
    assert inserted == 2

    pc1 = await db.get(PipeClass, "A1")
    pc2 = await db.get(PipeClass, "A2")
    assert pc1 is not None and pc2 is not None
    assert pc1.status == "DRAFT"
    assert pc1.flange_class == "PN25"


async def test_commit_import_rejects_when_errors(db):
    """preview 含 ERROR → commit 422 IMPORT_HAS_ERRORS, DB 无新行。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row(design_pressure=43.0)],
        sheet2=_valid_sheet2(),
    )
    svc = PipeClassImportService(db)
    _, import_id = await svc.preview(xlsx_bytes)
    with pytest.raises(PcsError) as ei:
        await svc.commit_import(import_id)
    assert ei.value.code == "IMPORT_HAS_ERRORS"
    assert ei.value.status == 422
    # DB 无新行（rollback）
    pc = await db.get(PipeClass, "A1")
    assert pc is None


async def test_commit_import_rejects_double_consume(db):
    """commit 一次成功，第二次同 import_id → 409 IMPORT_CONSUMED。"""
    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row()],
        sheet2=_valid_sheet2(),
    )
    svc = PipeClassImportService(db)
    _, import_id = await svc.preview(xlsx_bytes)
    inserted = await svc.commit_import(import_id)
    assert inserted == 1

    with pytest.raises(PcsError) as ei:
        await svc.commit_import(import_id)
    assert ei.value.code == "IMPORT_CONSUMED"
    assert ei.value.status == 409


# ---------------------------------------------------------------------------
# 12. 验证引擎集成
# ---------------------------------------------------------------------------


async def test_import_validation_engine_used(db, monkeypatch):
    """preview 实际调用 PipeClassValidator.validate_company（22 条规则）。"""
    calls: list[dict] = []

    real_validate = PipeClassValidator.validate_company

    def spy_validate(data, **kw):
        calls.append(data)
        return real_validate(data, **kw)

    monkeypatch.setattr(PipeClassValidator, "validate_company", staticmethod(spy_validate))

    xlsx_bytes = _xlsx(
        rows=[_make_sheet1_row(class_id="A1")],
        sheet2=_valid_sheet2("A1"),
    )
    svc = PipeClassImportService(db)
    await svc.preview(xlsx_bytes)
    assert calls, "preview 未触发 validator"
    assert calls[0]["class_id"] == "A1"
    assert "dn_series_json" in calls[0]
    assert "sch_series_json" in calls[0]
