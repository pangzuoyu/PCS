"""ExportService 测试（P2 Sprint 3 Task 5.3）。

测试覆盖：
1. 正常数据导出（3 测 + 1 perf budget，源自 brief Step 4/4.5）
2. status-report 数据导出端到端验证（Task 5.1 ↔ 5.3 集成）
"""

from __future__ import annotations

import resource
import time
from io import BytesIO

import openpyxl

from app.services.export_service import ExportService

# ---------------------------------------------------------------------------
# 基础导出测试（brief Step 1-4）
# ---------------------------------------------------------------------------


async def test_to_excel_produces_valid_workbook():
    data = [{"a": 1, "b": 2}, {"a": 3, "b": 4}]
    buf = await ExportService.to_excel(data, headers=["a", "b"])
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("a", "b")
    assert rows[1] == (1, 2)
    assert rows[2] == (3, 4)


async def test_to_excel_empty_data_only_header():
    buf = await ExportService.to_excel([], headers=["a"])
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows == [("a",)]


async def test_to_excel_custom_sheet_name():
    buf = await ExportService.to_excel([{"x": 1}], headers=["x"], sheet_name="P&ID")
    wb = openpyxl.load_workbook(buf)
    assert "P&ID" in wb.sheetnames


# ---------------------------------------------------------------------------
# 性能预算测试（brief Step 4.5 — D28 性能预算：10k 行 × 20 列 ≤ 2s + 内存 ≤ 50MB）
# ---------------------------------------------------------------------------


async def test_export_perf_budget():
    rows = [{f"col_{j}": f"v_{i}_{j}" for j in range(20)} for i in range(10_000)]
    # 不用 tracemalloc：tracemalloc hook Python 分配路径，4× 减速会把
    # write_only 实现从 0.9s 拖到 4s，掩盖真实瓶颈。改用 ``resource.getrusage``
    # 读 ru_maxrss（Linux 内核态高水位，单次 syscall，无 Python 侧开销）。
    start_rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss  # KiB
    start = time.monotonic()
    output = await ExportService.to_excel(
        rows, headers=[f"col_{j}" for j in range(20)]
    )
    elapsed = time.monotonic() - start
    peak_kib = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss - start_rss
    assert elapsed < 2.0, f"导出超时 {elapsed:.2f}s"
    assert peak_kib < 50 * 1024, f"导出内存增长 {peak_kib / 1024:.1f}MB"
    assert isinstance(output, BytesIO) and len(output.getvalue()) > 0


# ---------------------------------------------------------------------------
# 端到端集成测试：ReportService 数据 → Excel
# ---------------------------------------------------------------------------


async def test_to_excel_with_report_service_data(db, sample_assets_and_versions):
    """Task 5.1 status-report 数据可正确导出。

    验证 6 列对齐 + 值匹配。
    """
    from app.services.report_service import ReportService

    report = await ReportService(db).config_asset_status()
    data = [r.model_dump() for r in report]
    headers = ["category", "draft", "pending", "approved", "published", "obsolete"]
    buf = await ExportService.to_excel(data, headers=headers)
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    assert rows[0] == ("category", "draft", "pending", "approved", "published", "obsolete")
    assert len(rows) == 1 + len(data)  # header + N data rows
    for excel_row, source_obj in zip(rows[1:], report, strict=True):
        assert excel_row == (
            source_obj.category,
            source_obj.draft,
            source_obj.pending,
            source_obj.approved,
            source_obj.published,
            source_obj.obsolete,
        )
